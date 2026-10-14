from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s = doc.styles[name]; s.font.name = 'Calibri'; s.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
header = sec.header.paragraphs[0]
header.text = 'BEACON  /  SYSTEM DOCUMENTATION'
header.style = 'Caption'; header.runs[0].font.color.rgb = RGBColor.from_string('52636A')
footer = sec.footer.paragraphs[0]
footer.text = '12 October 2026  •  Configuration and code review     |     '
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.runs[0].font.size = Pt(9)
doc.core_properties.title = 'Beacon Technology Architecture and Cost'
doc.core_properties.author = 'Beacon Project Team'
doc.core_properties.subject = 'Current system flow, model transition and estimated operating cost'

def p(t, style=None): return doc.add_paragraph(t, style)
def h(t): doc.add_heading(t, 2)
def page(t): doc.add_page_break(); doc.add_heading(t, 1)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    widths=widths or {2:[1.8,5.2],3:[1.45,2,3.55],4:[2.1,1.35,1.35,2.2]}[len(headers)]
    t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr(); avoid=OxmlElement('w:cantSplit'); trPr.append(avoid)
        for j,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[j])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3); para.paragraph_format.space_before=Pt(2)
                for r in para.runs:r.font.size=Pt(10)
            if i==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7F0EE'); c._tc.get_or_add_tcPr().append(shade)
                for r in c.paragraphs[0].runs:r.bold=True
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t
def link(label,url):
    para=doc.add_paragraph(); hyperlink=OxmlElement('w:hyperlink')
    rid=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink.set(qn('r:id'),rid); run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    col=OxmlElement('w:color');col.set(qn('w:val'),'166A60');props.append(col);run.append(props)
    text=OxmlElement('w:t');text.text=label;run.append(text);hyperlink.append(run);para._p.append(hyperlink)

p('Beacon Technology Architecture and Cost','Title')
p('Management and implementation reference  |  12 October 2026','Subtitle')
h('What Beacon does')
p('Beacon is a conversational product guide for Leadrat CRM. A visitor can ask a product question, request a live screen demonstration and share business requirements. Beacon combines product documentation, conversation history and supported CRM actions to explain the product and prepare a structured sales handoff.')
h('Current model decision')
p('We are currently using the hosted DeepSeek Flash model for conversation planning, product answers, answer checking and customer qualification. The configured API model name is deepseek-flash. DeepSeek identifies this model as DeepSeek V4.1 Flash in its published model documentation [1]. Thinking mode is disabled by default in our integration to reduce response overhead.')
p('The earlier Qwen2.5 1.5B model was small and did not handle conversation context consistently in our tested setup. Examples included losing the meaning of short follow-up answers and selecting an unrelated product area. We therefore moved beyond that setup. A Qwen3 4B fine-tuning experiment followed, and DeepSeek is now the temporary application provider while we continue evaluating quality and speed.')
p('The Qwen adapter is not loaded into DeepSeek. DeepSeek runs on its provider’s servers through an API. The local Qwen3 release remains recorded in the project, but the current provider setting selects DeepSeek rather than that local model.')
h('What this document confirms')
table(['Area','Current position'],[
('Conversation model','DeepSeek Flash through the official DeepSeek API; no OpenAI model selected.'),
('Product knowledge','Handbook passages and CRM screen metadata are retrieved for each relevant request.'),
('Demo control','The model chooses from known screen options; browser code performs and checks the action.'),
('Database','SQLite for the current local setup; PostgreSQL with pgvector is supported for deployment.'),
('Cost and readiness','API estimates are provided on page 5. This document is not a new accuracy benchmark or production sign-off.')])
p('Verification basis: the current project configuration and source code were reviewed for this document. No new paid model benchmark or live CRM test was run during documentation preparation.')

page('Technology stack')
p('Each component has a separate responsibility. DeepSeek handles language decisions; the application retains control of storage, permitted actions and execution checks.')
table(['Component','Technology','Role in Beacon'],[
('Website and widget','HTML, CSS and JavaScript','Displays conversation, live CRM preview, session controls and voice options.'),
('Application server','Python, FastAPI and Uvicorn','Receives questions, manages sessions and coordinates the response and demonstration.'),
('Conversation workflow','LangGraph','Runs the sequence: load context, plan the turn, ground the answer and save the decision.'),
('Language model','DeepSeek Flash API with HTTPX','Understands the turn in context, chooses a screen, drafts answers and extracts customer details.'),
('Response validation','Pydantic and JSON contracts','Checks the structure and allowed values of model decisions before the app uses them.'),
('Document processing','pypdf','Reads the product handbook and divides it into searchable passages.'),
('Meaning search','Multilingual MiniLM, ONNX Runtime and NumPy','Creates compact representations of text locally and compares their meaning.'),
('Keyword and spelling search','BM25 and character similarity','Complements meaning search with exact product terms and spelling variations.'),
('Conversation storage','SQLite; optional PostgreSQL and pgvector','Stores session context, extracted facts, document passages and search representations.'),
('Live CRM demonstration','Playwright with a CRM screen map','Opens known pages and runs supported walkthroughs, then checks the displayed result.'),
('Voice input and output','Browser SpeechRecognition and edge-tts','Converts speech to text and reads replies aloud; browser speech synthesis is a fallback.'),
('Sales integration','JSON handoff, HTTP webhook and local outbox','Prepares the sales brief, records delivery attempts and prevents duplicate submissions.')])
h('Training tools belong to the earlier local model work')
p('Kaggle, PyTorch, Transformers, PEFT LoRA and bitsandbytes QLoRA were used for local Qwen fine-tuning. The recorded Qwen3 experiment used approximately 17,000 examples. llama.cpp and GGUF support local inference. These are retained project capabilities, not prerequisites for a DeepSeek API response. Whisper small is not the current speech recognition implementation.')

page('How a user question moves through Beacon')
p('The first destination is the Beacon application server, not the model provider. The server adds the context and product evidence needed to make the request useful.')
table(['Step','What happens'],[
('1  Visitor sends a question','The widget sends text to the session turn endpoint. Voice first becomes text through the browser speech recognition service.'),
('2  Server checks the session','FastAPI verifies session ownership, rejects empty or duplicate requests and starts the turn. Form confirmation and restricted write requests have separate checks.'),
('3  Conversation context is loaded','LangGraph assembles recent messages, accepted customer facts, the pending question, the last demonstrated feature and the last offer.'),
('4  Relevant evidence is found','Local search retrieves handbook passages and candidate CRM screens. Meaning, product words and spelling similarity help find candidates.'),
('5  DeepSeek plans the response','The model receives the question, context, evidence and allowed options. It returns a structured decision: product question, customer answer, clarification or another supported turn type, plus any proposed screen and fact updates.'),
('6  The decision is validated','The application checks the JSON and allowed values. An invalid or unclear result produces a clarification or retry message instead of an unchecked action.'),
('7  Answer and demo progress','For product questions, DeepSeek writes an evidence-based answer and checks its claims in another call. A selected screen can begin opening in parallel while the answer is prepared.'),
('8  The result reaches the visitor','The widget receives text, action status and updated preview. Voice reads the answer if enabled. A screen is reported as verified only after the browser action succeeds.'),
('9  Context and sales facts are updated','The decision is saved. New accepted customer facts can trigger a separate qualification call and a consent-controlled sales handoff.')])
h('A typical conversation')
p('If Beacon asks for monthly lead volume and the visitor replies “20k”, the pending question tells the planner that the reply is a customer fact. The accepted value can become 20,000 leads per month. If the visitor then asks “show me how to manage those leads”, the planner receives that history and searches for the relevant Leads workflow. This illustrates the intended flow; it is not a claim that every phrasing has passed a fresh test.')

page('Demo selection and customer data')
h('How the correct screen is chosen')
p('The system searches a catalogue built from CRM pages and sections, together with prepared walkthroughs. DeepSeek chooses among these options using the current question and conversation context. The same validated decision is used to guide the answer and start the demonstration, reducing disagreement between what Beacon says and what it shows.')
p('Demo intent is not decided solely by a list of exact user phrases or regex matches. However, the system still contains keyword search, spelling logic, validation rules and explicit browser instructions. Those are needed to locate screens and safely execute actions. It would be inaccurate to describe the entire project as having no rules or no hardcoded code.')
p('Playwright opens the selected CRM page or follows a supported walkthrough. A request to create a lead uses a separate form workflow with required values and confirmation. Opening an Add Lead form is not the same as saving a lead. Unsupported, unverified or blocked actions must not be presented as completed.')
h('What the sales team receives')
table(['Information group','Examples of values in the handoff'],[
('Customer and business','Name or contact details when supplied, company, business type, team size, role and geography.'),
('Business requirements','Monthly lead volume, current tools, lead sources, problems and product areas demonstrated.'),
('Qualification','Fit score and rationale, route, recommended next step and missing information.'),
('Evidence and permission','Conversation turn references, contact consent, timestamps and delivery status.')])
p('Unknown details remain empty rather than being invented. The model extracts information, while application rules validate and score it. If model extraction fails, the fallback normally sends the case for human review. Handoff requires the sales route, explicit permission and a usable contact method.')
h('Where information is stored and sent')
p('The current setup uses SQLite in the local application state folder. Conversation state and document search data are stored separately. PostgreSQL with pgvector is implemented as a deployment option but is not selected by the reviewed local configuration. Handoff records are kept in a local outbox; the default receiver is a mock CRM webhook, so a production sales endpoint still needs deployment configuration.')
p('DeepSeek receives the text included in model requests, such as recent conversation, relevant customer facts and selected product evidence. Browser speech recognition and edge-tts can also use external services. This is not a fully offline setup. Conversation and outbox cleanup default to seven days when their cleanup routines run; backups and other logs require their own retention policy.')

page('DeepSeek API cost')
p('Prices checked on 12 October 2026 against DeepSeek’s official pricing page [1]. Tokens are the small text units used for billing. Input includes instructions, history and retrieved evidence; output includes model-generated JSON and text, not only the visible answer.')
table(['Per 1 million tokens','Off peak','Peak'],[
('Input with no cache reuse','$0.15','$0.30'),
('Input reused from provider cache','$0.003','$0.006'),
('Output','$0.60','$1.20')])
p('Published peak hours are Monday to Friday, 01:00–04:00 and 06:00–10:00 UTC, equivalent to 06:30–09:30 and 11:30–15:30 IST. Other hours use off-peak rates. Prices can change; the provider’s billing record is authoritative.')
h('Budget for 20 30 and 50 question answer turns')
p('Planning assumption: 12,000 input tokens and 1,000 output tokens in total per user turn, aggregated across all model calls. This is an illustrative allowance, not measured Beacon usage. The estimate assumes no input cache savings and thinking disabled. INR uses an illustrative USD 1 = INR 90, not a verified exchange rate.')
table(['Questions with answers','Off peak USD','Peak USD','Approx INR range'],[
('20','$0.048','$0.096','4.32–8.64'),
('30','$0.072','$0.144','6.48–12.96'),
('50','$0.120','$0.240','10.80–21.60')])
p('Calculation per turn at peak: (12,000 ÷ 1,000,000 × $0.30) + (1,000 ÷ 1,000,000 × $1.20) = $0.0048. Off peak, the same workload costs $0.0024.')
h('Why actual cost varies')
p('A normal product turn commonly makes three calls: planning, answer generation and answer checking. A failed check can add a rewrite and another check. Customer qualification and network retries can add calls too. A simple customer answer can require fewer calls. Long histories and larger evidence selections increase input tokens. At twice the assumed token volume, the estimates double.')
p('For example, 1,000 conversations with 30 turns each would cost approximately $72–$144 in model API charges under the same assumptions, or INR 6,480–12,960 at the illustrative conversion. This is a usage scenario, not a monthly commitment or invoice.')
h('Costs excluded from the table')
p('Application hosting, browser automation capacity, database hosting and backups, monitoring, storage, domain, network traffic, speech services, taxes and payment conversion charges are separate. No deployment size or provider quote has been selected, so a defensible all-inclusive monthly total is not yet available. Actual API usage must be measured before approving a production budget.')

page('Performance and deployment status')
h('How the current design reduces waiting')
p('DeepSeek mode avoids loading the local Qwen language model onto the laptop for each application startup. Document search still runs locally. Search data is prepared in the background, embeddings and screen representations are cached, and a selected CRM screen can start opening before answer generation finishes. The code also supports preparing a browser in advance when that option is enabled.')
p('These changes reduce sources of delay, but do not make responses instant. Network latency, multiple model calls, CRM sign-in, page rendering, voice generation and retries still take time. This review did not establish a new measured response-time target or a DeepSeek accuracy percentage.')
h('What still needs production validation')
table(['Area','Required evidence before rollout'],[
('Conversation quality','Test follow-ups, corrections, interruptions, spelling errors, Hinglish, broad feature requests and missing context against expected outcomes.'),
('Answer and demonstration','Record whether both the answer and the displayed screen match the request. Verify saved records separately from opened forms.'),
('Speed and reliability','Measure first preview, screen change and full answer times; include slow API, timeout and CRM sign-in failures.'),
('Capacity','The reviewed application allows one active demo session. Multiple simultaneous visitors need further session and browser capacity work.'),
('Sales integration','Replace the mock receiver with the intended sales endpoint and test consent, duplicate delivery, failures and data retention.'),
('Cost tracking','Capture model input, cached input and output usage for every call. Current response wrappers do not provide a complete persisted token billing ledger.')])
h('Adapting Beacon to another product')
p('The conversation workflow and model connector can be reused. A new product still needs its own verified documentation, CRM screen map, supported walkthroughs, login configuration, business qualification rules and sales payload mapping. Updating knowledge alone will not automatically teach the browser how to perform every action in a different application.')
h('Current conclusion')
p('DeepSeek is the current temporary language-model provider. It replaces the earlier local-model path for conversation and qualification, while the application continues to control evidence retrieval, browser execution and customer data. The change addresses limitations observed in the earlier small-model setup, but production readiness must be demonstrated through fresh end-to-end tests rather than inferred from the model name.')

page('Verification references')
h('External sources')
link('[1] DeepSeek official models and pricing', 'https://api-docs.deepseek.com/quick_start/pricing/')
p('Used for the deepseek-flash model identity, peak and off-peak token prices, and billing schedule. Pricing checked on 12 October 2026. All cost scenarios in this document are our calculations from the stated assumptions.')
link('[2] DeepSeek official thinking mode documentation', 'https://api-docs.deepseek.com/guides/thinking_mode/')
p('Used to verify the supported thinking-mode switch. The application sends thinking disabled unless its explicit override is enabled.')
h('Project evidence reviewed')
table(['Source within the Beacon project','What it establishes'],[
('.env and app/config.py','Configured DeepSeek provider and model; voice provider. Secrets are intentionally excluded from this document.'),
('app/main.py and web/widget-client.js','Session entry point, turn execution, live preview, speech integration and single-demo capacity.'),
('app/conversation_engine.py','Actual LangGraph workflow, contextual planner, validation and saved state.'),
('app/local_model.py','Hosted DeepSeek endpoint, JSON output, thinking setting and transient retries.'),
('app/semantic_knowledge.py and app/embeddings.py','Handbook retrieval, local embeddings, answer generation and evidence checking.'),
('app/screens.py and knowledge/screens.json','Candidate screen search and CRM page metadata.'),
('app/conversation_store.py','SQLite storage and optional PostgreSQL with pgvector.'),
('app/qualify.py and app/handoff.py','Customer extraction, scoring, consent checks, JSON handoff and outbox delivery.'),
('slm/current_release.json and requirements.txt','Retained Qwen3 release record and installed application dependencies.')])
h('How to interpret the model history')
p('The earlier Qwen2.5 1.5B limitations described here come from the project’s previously observed conversations and feedback. They are not a controlled comparison proving that model size alone caused every issue. Retrieval quality, instructions, conversation state and browser execution also affect the final result.')
p('No historical Qwen test percentage has been reused as a DeepSeek success rate.')

path=OUT/'Beacon_Technology_Architecture_and_Cost.docx'
doc.save(path)
print(path)
